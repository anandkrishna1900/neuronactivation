"""Convert Phase 6B markdown report to Word document."""
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parent.parent
md_path = ROOT / "docs" / "FlyMind_Phase6B_Motor_Readout_Report.md"
docx_path = ROOT / "docs" / "FlyMind_Phase6B_Motor_Readout_Report.docx"

with open(md_path, "r", encoding="utf-8") as f:
    lines = f.readlines()

doc = Document()

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(4)

for line in lines:
    line = line.rstrip("\n")

    if line.startswith("# "):
        doc.add_heading(line[2:], level=0)
    elif line.startswith("## "):
        doc.add_heading(line[3:], level=1)
    elif line.startswith("### "):
        doc.add_heading(line[4:], level=2)
    elif line.startswith("---"):
        doc.add_paragraph("─" * 60)
    elif line.startswith("| ") and "---" not in line:
        cells = [c.strip() for c in line.split("|")[1:-1]]
        if not hasattr(doc, "_last_table") or doc._last_table is None:
            doc._last_table = doc.add_table(rows=1, cols=len(cells))
            doc._last_table.style = "Table Grid"
            for i, cell in enumerate(cells):
                doc._last_table.rows[0].cells[i].text = cell
        else:
            row = doc._last_table.add_row()
            for i, cell in enumerate(cells):
                row.cells[i].text = cell
    elif line.startswith("| ") and "---" in line:
        doc._last_table = None
    elif line.startswith("- "):
        doc.add_paragraph(line[2:], style="List Bullet")
    elif line.startswith("    ") and line.strip():
        p = doc.add_paragraph()
        run = p.add_run(line.strip())
        run.font.name = "Consolas"
        run.font.size = Pt(10)
    elif line.strip():
        doc.add_paragraph(line)
    else:
        doc._last_table = None

doc.save(str(docx_path))
print("Saved: %s" % docx_path)
